"""The shipped detector: calibration, verdicts, warnings, files, CLI, app, privacy."""

import json
import os
import socket
from pathlib import Path

import numpy as np
import pytest
from hypothesis import given, settings
from hypothesis import strategies as st
from sklearn.isotonic import IsotonicRegression

from aidetect.calibration import VennAbers, _pava
from aidetect.cli import main as cli_main
from aidetect.detector import DISPLAY_MAX, DISPLAY_MIN, Detector, classify
from aidetect.ingest import UnsupportedFileError, read_file
from aidetect.reliability import check
from aidetect.stylometry import FEATURES, extract

HUMAN = (
    "I wasn't going to write about the garden this year, honestly. Last spring the slugs got "
    "everything, the beans, the lettuce, even the marigolds my neighbour swore they wouldn't "
    "touch. So this year I just... let it go. Grass up to my knees by June. And then something "
    "odd happened: poppies. Hundreds of them, from seed that must have been sitting there since "
    "my grandmother's time. My kids think I planted them on purpose. I haven't told them.\n\n"
    "We had people over on Saturday and everyone asked about it. Dave (who knows a lot about "
    "this sort of thing, or says he does) reckons the soil got disturbed when they dug up the "
    "drains. Maybe. I prefer the grandmother theory. Anyway, the slugs are back, but they don't "
    "seem to like poppies, which is the first piece of luck I've had with that bed in a decade. "
    "Next year I might actually plant something deliberately. Or I might not."
)
AI = (
    "Gardening offers numerous benefits for both physical and mental well-being. Firstly, it "
    "provides an excellent opportunity for moderate physical activity, which can improve "
    "cardiovascular health and overall fitness. Furthermore, spending time outdoors in nature "
    "has been shown to reduce stress and enhance mood. Additionally, growing one's own fruits "
    "and vegetables promotes healthier eating habits and fosters a deeper connection to food.\n\n"
    "Moreover, gardening plays a crucial role in environmental sustainability. By cultivating "
    "native plants, gardeners can support local ecosystems and provide habitats for pollinators "
    "such as bees and butterflies. It is important to note that composting and responsible "
    "water use further contribute to a healthier planet. In conclusion, gardening is a "
    "rewarding activity that nurtures both the individual and the wider community, making it a "
    "valuable pursuit for people of all ages and backgrounds."
)


@pytest.fixture(scope="module")
def detector() -> Detector:
    return Detector()


class TestCalibration:
    def test_pava_is_monotone_least_squares(self) -> None:
        rng = np.random.default_rng(0)
        y = rng.random(200)
        fit = _pava(y)
        assert np.all(np.diff(fit) >= -1e-12)
        ref = IsotonicRegression().fit(np.arange(200), y).predict(np.arange(200))
        np.testing.assert_allclose(fit, ref, atol=1e-9)

    @settings(max_examples=50, deadline=None)
    @given(st.floats(-6, 6))
    def test_interval_contains_point(self, score: float) -> None:
        rng = np.random.default_rng(1)
        scores = rng.normal(size=300)
        labels = (scores + rng.normal(scale=0.8, size=300) > 0).astype(float)
        p, low, high = VennAbers(scores, labels).predict(score)
        assert 0.0 <= low <= p <= high <= 1.0

    def test_round_trip(self) -> None:
        va = VennAbers(np.array([0.1, 0.5, 0.9]), np.array([0, 1, 1]))
        again = VennAbers.from_dict(va.to_dict())
        assert again.predict(0.6) == va.predict(0.6)


class TestVerdicts:
    def test_bands(self) -> None:
        assert classify(0.99, 1.0, 0.98, blocking=False) == "likely_ai"
        assert classify(0.99, 1.0, 0.98, blocking=True) == "inconclusive"
        assert classify(0.10, 0.20, 0.98, blocking=False) == "likely_human"
        assert classify(0.60, 0.70, 0.98, blocking=False) == "inconclusive"

    def test_never_certain(self, detector: Detector) -> None:
        for text in (HUMAN, AI, AI * 4, HUMAN * 4):
            r = detector.analyse(text)
            assert r.probability is not None and r.low is not None and r.high is not None
            assert DISPLAY_MIN <= r.low <= r.probability <= r.high <= DISPLAY_MAX

    def test_examples_rank_sensibly(self, detector: Detector) -> None:
        human, ai = detector.analyse(HUMAN), detector.analyse(AI)
        assert human.probability is not None and ai.probability is not None
        assert ai.probability > human.probability
        assert human.band != "likely_ai"

    def test_short_text_gets_no_verdict(self, detector: Detector) -> None:
        r = detector.analyse("This is far too short to judge.")
        assert r.band == "inconclusive"
        assert any(w.code == "short" for w in r.warnings)

    def test_empty(self, detector: Detector) -> None:
        r = detector.analyse("   \n ")
        assert r.probability is None and r.n_words == 0

    def test_result_is_json_and_heatmap_is_relative(self, detector: Detector) -> None:
        r = detector.analyse(HUMAN + "\n\n" + AI)
        data = json.loads(json.dumps(r.to_dict()))
        assert data["verdict"] in {"Likely AI-generated", "Inconclusive", "Likely human-written"}
        heats = [s.heat for s in r.sentences]
        assert heats and min(heats) == 0.0 and max(heats) <= 1.0
        faint = detector.analyse(HUMAN)
        assert max(s.heat for s in faint.sentences) < 0.2
        assert r.reasons_ai or r.reasons_human

    def test_analysis_never_uses_the_network(
        self, detector: Detector, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        def refuse(*_args: object, **_kwargs: object) -> None:
            raise AssertionError("network access during analysis")

        monkeypatch.setattr(socket.socket, "connect", refuse)
        monkeypatch.setattr(socket, "create_connection", refuse)
        assert detector.analyse(AI).probability is not None


class TestReliability:
    def codes(self, text: str) -> set[str]:
        return {w.code for w in check(text)}

    def test_code_lists_quotes_language(self) -> None:
        code = "\n".join(["def f(x):", "    return x == 1", "if (a) {", "  b();", "}"] * 30)
        assert "code" in self.codes(code)
        bullets = "\n".join(f"- item number {i} with a few words" for i in range(40))
        assert "lists" in self.codes(bullets)
        quoted = 'He said "' + "word " * 200 + '" and left.'
        assert "quotes" in self.codes(quoted)
        german = "Das ist ein ganz normaler deutscher Text über das Wetter und die Stadt. " * 20
        assert "language" in self.codes(german)
        assert "short" not in self.codes(HUMAN + " " + HUMAN)


def test_stylometry_features_are_finite() -> None:
    for text in (HUMAN, AI, "One.", "a " * 500):
        feats = extract(text)
        assert set(feats) == set(FEATURES)
        assert all(np.isfinite(v) for v in feats.values())
    assert extract(AI)["transition_rate"] > extract(HUMAN)["transition_rate"]
    assert extract(HUMAN)["contraction_rate"] > extract(AI)["contraction_rate"]


class TestIngest:
    def test_txt_encodings(self, tmp_path: Path) -> None:
        (tmp_path / "a.txt").write_bytes("café naïve".encode("utf-16"))
        assert read_file(tmp_path / "a.txt") == "café naïve"
        (tmp_path / "b.txt").write_bytes("caf\xe9".encode("cp1252"))
        assert read_file(tmp_path / "b.txt") == "café"

    def test_docx(self, tmp_path: Path) -> None:
        import docx

        doc = docx.Document()
        doc.add_paragraph("First paragraph.")
        doc.add_paragraph("Second paragraph.")
        doc.save(str(tmp_path / "t.docx"))
        assert read_file(tmp_path / "t.docx") == "First paragraph.\n\nSecond paragraph."

    def test_unsupported(self, tmp_path: Path) -> None:
        (tmp_path / "x.xlsx").write_bytes(b"x")
        with pytest.raises(UnsupportedFileError):
            read_file(tmp_path / "x.xlsx")


class TestCli:
    def test_text_and_json(self, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
        path = tmp_path / "essay.txt"
        path.write_text(AI)
        assert cli_main([str(path), "--no-color", "--heatmap"]) == 0
        out = capsys.readouterr().out
        assert "AI-likeness" in out and "not proof" in out
        assert cli_main([str(path), "--json"]) == 0
        assert "probability_ai" in json.loads(capsys.readouterr().out)

    def test_unreadable_and_short(self, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
        assert cli_main([str(tmp_path / "missing.txt")]) == 2
        short = tmp_path / "short.txt"
        short.write_text("")
        assert cli_main([str(short), "--json"]) == 3
        capsys.readouterr()


def test_app_window_renders_offscreen(detector: Detector) -> None:
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    pytest.importorskip("PySide6")
    from PySide6.QtWidgets import QApplication

    from aidetect.app.main import MainWindow, _newer, about_html, render_report

    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    result = detector.analyse(AI)
    window.show_result(result)
    assert window.verdict.text() == result.band_label
    page = render_report(result)
    assert "not proof" in page
    assert "Measured accuracy" in about_html(detector, detector.metrics)
    assert _newer("0.2.0", "0.1.0") and not _newer("0.1.0", "0.1.0")
    window.worker_thread.quit()
    window.worker_thread.wait(5000)
    window.close()
    del app


def test_theme_is_readable_in_system_dark_mode() -> None:
    # Regression: under Windows dark mode Qt handed the app light text, which
    # was drawn on the app's light backgrounds and became invisible.
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    pytest.importorskip("PySide6")
    from PySide6.QtGui import QColor, QPalette
    from PySide6.QtWidgets import QApplication

    from aidetect.app.main import apply_theme

    app = QApplication.instance() or QApplication([])
    dark = QPalette()
    dark.setColor(QPalette.ColorRole.Text, QColor("#f0f0f0"))
    dark.setColor(QPalette.ColorRole.WindowText, QColor("#f0f0f0"))
    app.setPalette(dark)
    apply_theme(app)
    palette = app.palette()
    for role in (QPalette.ColorRole.Text, QPalette.ColorRole.WindowText):
        assert palette.color(role).lightness() < 80
    assert palette.color(QPalette.ColorRole.Base).lightness() > 200
